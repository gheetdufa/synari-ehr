"""Reading SimplePractice reports saved as CSV or Excel (.xlsx) into rows of text.

Reports may start with a title or date-range lines, so the heading row is the one among the
first rows that matches the most known column names. Blank rows and "Total" rows are skipped.
"""

import csv
import io
from collections.abc import Iterable, Iterator
from dataclasses import dataclass
from datetime import date, datetime, time
from pathlib import Path
from typing import BinaryIO
from zipfile import BadZipFile

from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException

from synari.services.imports import ImportFailed
from synari.services.imports.fields import KINDS, normalize_header

HEADER_SCAN_ROWS = 15
MAX_ROWS = 200_000
# Reports are read into memory; the 2 GB upload limit is meant for the export zip.
MAX_REPORT_BYTES = 200 * 1024 * 1024
MAX_CELL_CHARS = 2_000
HEADER_SNIFF_BYTES = 256 * 1024
_TOTAL_ROWS = frozenset({"total", "totals", "grand total"})
_KNOWN_HEADERS = frozenset(
    alias for definition in KINDS.values() for f in definition.fields for alias in f.aliases
)


@dataclass
class Table:
    headers: list[str]
    rows: list[list[str]]

    def records(self, mapping: dict[str, str | None]) -> list[dict[str, str]]:
        """Each row as {field: value} for the mapped fields."""
        positions = {
            field: self.headers.index(header)
            for field, header in mapping.items()
            if header and header in self.headers
        }
        return [
            {field: row[index] for field, index in positions.items() if index < len(row)}
            for row in self.rows
        ]


def _cell(value: object) -> str:
    if value is None:
        return ""
    if isinstance(value, datetime):
        if value.time() == time(0, 0):
            return value.date().isoformat()
        return value.strftime("%Y-%m-%d %H:%M")
    if isinstance(value, date):
        return value.isoformat()
    if isinstance(value, time):
        return value.strftime("%H:%M")
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    return str(value).strip()[:MAX_CELL_CHARS]


def _header_score(row: list[str]) -> int:
    return sum(1 for cell in row if normalize_header(cell) in _KNOWN_HEADERS)


def _find_header(rows: list[list[str]]) -> int:
    """Index of the heading row among the first rows (the first non-blank row by default)."""
    best, best_score = -1, 1
    for index, row in enumerate(rows[:HEADER_SCAN_ROWS]):
        score = _header_score(row)
        if score > best_score:
            best, best_score = index, score
    if best >= 0:
        return best
    return next((i for i, row in enumerate(rows) if any(row)), 0)


def _is_blank_or_total(row: list[str]) -> bool:
    filled = [cell for cell in row if cell]
    return not filled or normalize_header(filled[0]) in _TOTAL_ROWS


def _build(rows: Iterable[list[str]], filename: str, *, headers_only: bool = False) -> Table:
    head: list[list[str]] = []
    iterator = iter(rows)
    for row in iterator:
        head.append(row)
        if len(head) >= HEADER_SCAN_ROWS:
            break
    if not any(any(row) for row in head):
        raise ImportFailed(f"{filename} is empty")
    header_index = _find_header(head)
    headers = [cell.strip() for cell in head[header_index]]
    while headers and not headers[-1]:
        headers.pop()
    headers = [h or f"Column {i + 1}" for i, h in enumerate(headers)]
    if headers_only:
        return Table(headers, [])
    data: list[list[str]] = []
    for row in [*head[header_index + 1 :], *iterator]:
        if _is_blank_or_total(row):
            continue
        data.append(row[: len(headers)] + [""] * (len(headers) - len(row)))
        if len(data) > MAX_ROWS:
            raise ImportFailed(f"{filename} has more than {MAX_ROWS:,} rows; split it up")
    return Table(headers, data)


def decode_text(raw: bytes) -> str:
    """Text from a CSV or TXT file: UTF-8 (with or without BOM), else Windows-1252."""
    for encoding in ("utf-8-sig", "cp1252"):
        try:
            return raw.decode(encoding)
        except UnicodeDecodeError:
            continue
    return raw.decode("latin-1")


def _csv_rows(text: str) -> Iterator[list[str]]:
    first_line = text.split("\n", 1)[0]
    delimiter = max((",", ";", "\t"), key=first_line.count)
    for row in csv.reader(io.StringIO(text), delimiter=delimiter):
        yield [cell.strip()[:MAX_CELL_CHARS] for cell in row]


def _xlsx_rows(source: Path | BinaryIO, filename: str) -> Iterator[list[str]]:
    try:
        workbook = load_workbook(source, read_only=True, data_only=True)
    except (InvalidFileException, BadZipFile, OSError, KeyError, ValueError) as exc:
        raise ImportFailed(
            f"{filename} couldn't be opened as an Excel workbook; save it as .xlsx or CSV"
        ) from exc
    try:
        sheet = workbook.worksheets[0]
        for values in sheet.iter_rows(values_only=True):
            yield [_cell(value) for value in values]
    finally:
        workbook.close()


def csv_table(text: str, filename: str) -> Table:
    """A CSV report already decoded to text."""
    return _build(_csv_rows(text), filename)


def read_table(path: Path, filename: str) -> Table:
    """Read a whole CSV or XLSX report."""
    if path.stat().st_size > MAX_REPORT_BYTES:
        raise ImportFailed(f"{filename} is too large for a report; export a shorter date range")
    if filename.lower().endswith(".xlsx"):
        return _build(_xlsx_rows(path, filename), filename)
    return csv_table(decode_text(path.read_bytes()), filename)


def read_headers(upload: BinaryIO, filename: str) -> list[str]:
    """The heading row of an uploaded report (reads only the start of a CSV)."""
    upload.seek(0)
    try:
        if filename.lower().endswith(".xlsx"):
            return _build(_xlsx_rows(upload, filename), filename, headers_only=True).headers
        text = decode_text(upload.read(HEADER_SNIFF_BYTES))
        # The last line may be cut off mid-way; the headings are near the top anyway.
        return _build(
            _csv_rows(text.rsplit("\n", 1)[0] or text), filename, headers_only=True
        ).headers
    finally:
        upload.seek(0)
